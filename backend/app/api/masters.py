"""Master data management + OCR match/link endpoints.
Company admins manage types/records; all tenant members can view and link."""

import csv
import io
import uuid
from datetime import datetime

from fastapi import APIRouter, Depends, Form, HTTPException, UploadFile, status
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.api.deps import require_company_admin, require_company_member
from app.db.session import get_db
from app.events.publisher import publish_event
from app.models import (
    Document,
    MasterMatch,
    MasterRecord,
    MasterType,
    MatchStatus,
    Page,
    User,
)
from app.services import matching
from app.services.audit import log_action

router = APIRouter(prefix="/masters", tags=["masters"])
match_router = APIRouter(tags=["masters"])


# --- Schemas ---

class MasterFieldSchema(BaseModel):
    key: str
    label: str
    matchable: bool = False
    required: bool = False


class MasterTypeIn(BaseModel):
    name: str = Field(min_length=1, max_length=255)
    fields: list[MasterFieldSchema]


class MasterTypeOut(BaseModel):
    id: uuid.UUID
    name: str
    fields: list[dict]
    records_count: int = 0
    created_at: datetime

    model_config = {"from_attributes": True}


class MasterRecordIn(BaseModel):
    data: dict[str, str]


class MasterRecordOut(BaseModel):
    id: uuid.UUID
    master_type_id: uuid.UUID
    data: dict
    created_at: datetime

    model_config = {"from_attributes": True}


class MatchOut(BaseModel):
    id: uuid.UUID
    page_id: uuid.UUID
    master_record_id: uuid.UUID
    master_type_id: uuid.UUID
    master_type_name: str
    field_key: str
    field_label: str
    matched_text: str
    master_value: str
    record_data: dict
    field_labels: dict[str, str]  # key → human label, for the detail popup
    score: float
    kind: str
    status: str


# --- Helpers ---

def _get_type(db: Session, user: User, type_id: uuid.UUID) -> MasterType:
    master_type = db.get(MasterType, type_id)
    if master_type is None or master_type.company_id != user.company_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Master type not found")
    return master_type


def _get_record(db: Session, user: User, record_id: uuid.UUID) -> MasterRecord:
    record = db.get(MasterRecord, record_id)
    if record is None or record.company_id != user.company_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Master record not found")
    return record


def _validate_record_data(master_type: MasterType, data: dict) -> dict:
    keys = {field["key"] for field in master_type.fields}
    cleaned = {key: str(value).strip() for key, value in data.items() if key in keys}
    for field in master_type.fields:
        if field.get("required") and not cleaned.get(field["key"]):
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST, f"field '{field['label']}' is required"
            )
    return cleaned


# --- Master types ---

@router.get("/types", response_model=list[MasterTypeOut])
def list_types(user: User = Depends(require_company_member), db: Session = Depends(get_db)):
    counts = {
        row[0]: row[1]
        for row in db.execute(
            select(MasterRecord.master_type_id, func.count())
            .where(MasterRecord.company_id == user.company_id)
            .group_by(MasterRecord.master_type_id)
        )
    }
    types = db.scalars(
        select(MasterType)
        .where(MasterType.company_id == user.company_id)
        .order_by(MasterType.created_at)
    ).all()
    return [
        MasterTypeOut(
            id=t.id, name=t.name, fields=t.fields,
            records_count=counts.get(t.id, 0), created_at=t.created_at,
        )
        for t in types
    ]


@router.post("/types", response_model=MasterTypeOut, status_code=status.HTTP_201_CREATED)
def create_type(
    body: MasterTypeIn,
    admin: User = Depends(require_company_admin),
    db: Session = Depends(get_db),
):
    try:
        fields = matching.validate_fields([field.model_dump() for field in body.fields])
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc))
    if db.scalar(
        select(MasterType).where(
            MasterType.company_id == admin.company_id, MasterType.name == body.name
        )
    ):
        raise HTTPException(status.HTTP_409_CONFLICT, "A master type with this name exists")
    master_type = MasterType(
        company_id=admin.company_id, name=body.name, fields=fields, created_by=admin.id
    )
    db.add(master_type)
    db.flush()
    log_action(db, "master_type.create", company_id=admin.company_id, actor_user_id=admin.id,
               target_type="master_type", target_id=str(master_type.id))
    db.commit()
    publish_event(admin.company_id, "masters.changed", {"action": "type_created"})
    return MasterTypeOut(
        id=master_type.id, name=master_type.name, fields=master_type.fields,
        records_count=0, created_at=master_type.created_at,
    )


@router.patch("/types/{type_id}", response_model=MasterTypeOut)
def update_type(
    type_id: uuid.UUID,
    body: MasterTypeIn,
    admin: User = Depends(require_company_admin),
    db: Session = Depends(get_db),
):
    master_type = _get_type(db, admin, type_id)
    try:
        master_type.fields = matching.validate_fields(
            [field.model_dump() for field in body.fields]
        )
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc))
    master_type.name = body.name
    db.commit()
    publish_event(admin.company_id, "masters.changed", {"action": "type_updated"})
    records = db.scalar(
        select(func.count()).select_from(MasterRecord).where(
            MasterRecord.master_type_id == master_type.id
        )
    )
    return MasterTypeOut(
        id=master_type.id, name=master_type.name, fields=master_type.fields,
        records_count=records, created_at=master_type.created_at,
    )


@router.delete("/types/{type_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_type(
    type_id: uuid.UUID,
    admin: User = Depends(require_company_admin),
    db: Session = Depends(get_db),
):
    master_type = _get_type(db, admin, type_id)
    db.delete(master_type)  # records + matches cascade
    log_action(db, "master_type.delete", company_id=admin.company_id, actor_user_id=admin.id,
               target_type="master_type", target_id=str(type_id))
    db.commit()
    publish_event(admin.company_id, "masters.changed", {"action": "type_deleted"})


# --- Records ---

@router.get("/types/{type_id}/records")
def list_records(
    type_id: uuid.UUID,
    q: str | None = None,
    page: int = 1,
    page_size: int = 25,
    user: User = Depends(require_company_member),
    db: Session = Depends(get_db),
):
    master_type = _get_type(db, user, type_id)
    page = max(page, 1)
    page_size = min(max(page_size, 1), 100)
    records = db.scalars(
        select(MasterRecord)
        .where(MasterRecord.master_type_id == master_type.id)
        .order_by(MasterRecord.created_at.desc())
    ).all()
    if q:
        needle = q.lower()
        records = [
            record for record in records
            if any(needle in str(value).lower() for value in record.data.values())
        ]
    total = len(records)
    window = records[(page - 1) * page_size : page * page_size]
    return {
        "items": [MasterRecordOut.model_validate(r, from_attributes=True) for r in window],
        "total": total,
        "page": page,
        "page_size": page_size,
    }


@router.post(
    "/types/{type_id}/records",
    response_model=MasterRecordOut,
    status_code=status.HTTP_201_CREATED,
)
def create_record(
    type_id: uuid.UUID,
    body: MasterRecordIn,
    admin: User = Depends(require_company_admin),
    db: Session = Depends(get_db),
):
    master_type = _get_type(db, admin, type_id)
    record = MasterRecord(
        company_id=admin.company_id,
        master_type_id=master_type.id,
        data=_validate_record_data(master_type, body.data),
        created_by=admin.id,
    )
    db.add(record)
    db.commit()
    publish_event(admin.company_id, "masters.changed", {"action": "record_created"})
    return record


@router.patch("/records/{record_id}", response_model=MasterRecordOut)
def update_record(
    record_id: uuid.UUID,
    body: MasterRecordIn,
    admin: User = Depends(require_company_admin),
    db: Session = Depends(get_db),
):
    record = _get_record(db, admin, record_id)
    master_type = db.get(MasterType, record.master_type_id)
    record.data = _validate_record_data(master_type, body.data)
    db.commit()
    publish_event(admin.company_id, "masters.changed", {"action": "record_updated"})
    return record


@router.delete("/records/{record_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_record(
    record_id: uuid.UUID,
    admin: User = Depends(require_company_admin),
    db: Session = Depends(get_db),
):
    record = _get_record(db, admin, record_id)
    db.delete(record)
    db.commit()
    publish_event(admin.company_id, "masters.changed", {"action": "record_deleted"})


def _infer_fields(header: list[str]) -> list[dict]:
    """Build field definitions from an import file's header row, for master
    types created with no fields (name only). ASCII headers become the key
    directly; Japanese headers get a generated key with the text as label."""
    import re

    fields: list[dict] = []
    seen: set[str] = set()
    for index, cell in enumerate(header, start=1):
        label = cell.strip()
        if not label:
            continue
        if re.fullmatch(r"[A-Za-z0-9_\- ]+", label):
            key = re.sub(r"[^a-z0-9_]", "_", label.lower()).strip("_") or f"col_{index}"
        else:
            key = f"col_{index}"
        if key in seen:
            key = f"{key}_{index}"
        seen.add(key)
        fields.append({"key": key, "label": label, "matchable": True, "required": False})
    return fields


_NUMBERISH = None  # compiled lazily


def _is_numberish(value: str) -> bool:
    import re

    global _NUMBERISH
    if _NUMBERISH is None:
        _NUMBERISH = re.compile(r"[\d,.\-¥￥%/年月日: ]+")
    return bool(value) and _NUMBERISH.fullmatch(value) is not None


def _guess_has_header(rows: list[list[str]], master_type: MasterType) -> bool:
    """Heuristic: does the first row look like a header or like data?

    A plain product list without a header row must not lose its first item,
    so default to "data" unless something clearly marks row 1 as a header.
    """
    first = rows[0]
    # 1. A cell matches the type's field keys/labels → definitely a header.
    if master_type.fields:
        known = {field["key"] for field in master_type.fields} | {
            field["label"] for field in master_type.fields
        }
        if any(cell in known for cell in first):
            return True
    # 2. A column is numeric in the data rows but not in row 1 → header.
    sample = rows[1:9]
    if sample:
        for col, cell in enumerate(first):
            values = [row[col] for row in sample if col < len(row) and row[col]]
            if cell and values and not _is_numberish(cell) and all(
                _is_numberish(v) for v in values
            ):
                return True
    # 3. Row 1 uses typical header words (商品名, コード, 単価, name, code …).
    hints = (
        "名", "コード", "番号", "単価", "価格", "数量", "金額", "日付", "住所",
        "name", "code", "price", "amount", "qty", "id", "no.", "email", "tel",
    )
    if any(hint in cell.lower() for cell in first if cell for hint in hints):
        return True
    return False


def _read_tabular(filename: str, raw: bytes) -> list[list[str]]:
    """Rows (header first) from a .csv or .xlsx upload.

    CSVs from Japanese Excel are often Shift-JIS; try UTF-8 first and fall
    back to cp932 so 顧客名-style headers survive.
    """
    if filename.lower().endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST,
                "Excel import requires the 'openpyxl' package on the server",
            )
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        rows = [
            ["" if cell is None else str(cell).strip() for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        workbook.close()
        return rows
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp932", errors="replace")
    return [
        [cell.strip() for cell in row]
        for row in csv.reader(io.StringIO(text))
    ]


@router.post("/types/{type_id}/import/preview")
async def preview_import(
    type_id: uuid.UUID,
    file: UploadFile,
    admin: User = Depends(require_company_admin),
    db: Session = Depends(get_db),
):
    """First rows of the file + a guess whether row 1 is a header, so the UI
    can ask the user to confirm before anything is written."""
    master_type = _get_type(db, admin, type_id)
    raw = await file.read()
    rows = _read_tabular(file.filename or "upload.csv", raw)
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "File is empty")
    return {
        "rows": rows[:6],
        "total_rows": len(rows),
        "guessed_header": _guess_has_header(rows, master_type),
    }


@router.post("/types/{type_id}/import")
async def import_records(
    type_id: uuid.UUID,
    file: UploadFile,
    has_header: bool | None = Form(None),
    admin: User = Depends(require_company_admin),
    db: Session = Depends(get_db),
):
    """CSV / Excel bulk import. Header cells may be either the field keys or
    the field labels (顧客名 …). `has_header=false` imports every row as data
    (columns map to the type's fields in order). Duplicate rows are skipped."""
    master_type = _get_type(db, admin, type_id)
    raw = await file.read()
    rows = _read_tabular(file.filename or "upload.csv", raw)
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "File is empty")

    header_present = _guess_has_header(rows, master_type) if has_header is None else has_header

    # A type created with no fields gets its columns from the file.
    if not master_type.fields:
        if header_present:
            inferred = _infer_fields(rows[0])
        else:
            width = max(len(row) for row in rows)
            inferred = [
                {"key": f"col_{i}", "label": f"列{i}", "matchable": True, "required": False}
                for i in range(1, width + 1)
            ]
        if not inferred:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Header row is empty")
        master_type.fields = inferred
        db.flush()

    if header_present:
        # Map each header cell to a field key — by key or by label.
        known_keys = {field["key"] for field in master_type.fields}
        label_to_key = {field["label"]: field["key"] for field in master_type.fields}
        header = [label_to_key.get(cell, cell) for cell in rows[0]]
        if not any(key in known_keys for key in header):
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST,
                f"Header must use the field keys {sorted(known_keys)} "
                f"or labels {sorted(label_to_key)}",
            )
        data_rows = rows[1:]
        first_line = 2
    else:
        # No header — columns map to the type's fields in definition order.
        header = [field["key"] for field in master_type.fields]
        data_rows = rows
        first_line = 1

    existing = {
        tuple(sorted(record.data.items()))
        for record in db.scalars(
            select(MasterRecord).where(MasterRecord.master_type_id == master_type.id)
        )
    }
    created = 0
    skipped = 0
    errors: list[str] = []
    for line_number, cells in enumerate(data_rows, start=first_line):
        row = {key: value for key, value in zip(header, cells) if key}
        try:
            data = _validate_record_data(master_type, row)
        except HTTPException as exc:
            errors.append(f"line {line_number}: {exc.detail}")
            continue
        signature = tuple(sorted(data.items()))
        if not any(data.values()) or signature in existing:
            skipped += 1
            continue
        existing.add(signature)
        db.add(
            MasterRecord(
                company_id=admin.company_id,
                master_type_id=master_type.id,
                data=data,
                created_by=admin.id,
            )
        )
        created += 1
    log_action(db, "master.import_csv", company_id=admin.company_id, actor_user_id=admin.id,
               target_type="master_type", target_id=str(type_id),
               detail={"created": created, "skipped": skipped})
    db.commit()
    publish_event(admin.company_id, "masters.changed", {"action": "imported"})
    return {"created": created, "skipped": skipped, "errors": errors[:20]}


# --- Matches (mounted without the /masters prefix) ---

@match_router.get("/pages/{page_id}/matches", response_model=list[MatchOut])
def page_matches(
    page_id: uuid.UUID,
    user: User = Depends(require_company_member),
    db: Session = Depends(get_db),
):
    page = db.get(Page, page_id)
    if page is None or page.company_id != user.company_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Page not found")
    rows = db.execute(
        select(MasterMatch, MasterRecord, MasterType)
        .join(MasterRecord, MasterMatch.master_record_id == MasterRecord.id)
        .join(MasterType, MasterMatch.master_type_id == MasterType.id)
        .where(
            MasterMatch.page_id == page.id,
            MasterMatch.status != MatchStatus.dismissed.value,
        )
        .order_by(MasterMatch.score.desc())
    ).all()
    out = []
    for match, record, master_type in rows:
        labels = {field["key"]: field["label"] for field in master_type.fields}
        out.append(
            MatchOut(
                id=match.id,
                page_id=match.page_id,
                master_record_id=record.id,
                master_type_id=master_type.id,
                master_type_name=master_type.name,
                field_key=match.field_key,
                field_label=labels.get(match.field_key, match.field_key),
                matched_text=match.matched_text,
                master_value=match.master_value,
                record_data=record.data,
                field_labels=labels,
                score=match.score,
                kind=match.kind,
                status=match.status,
            )
        )
    return out


def _get_match(db: Session, user: User, match_id: uuid.UUID) -> MasterMatch:
    match = db.get(MasterMatch, match_id)
    if match is None or match.company_id != user.company_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Match not found")
    return match


@match_router.post("/matches/{match_id}/link")
def link(
    match_id: uuid.UUID,
    user: User = Depends(require_company_member),
    db: Session = Depends(get_db),
):
    match = _get_match(db, user, match_id)
    try:
        matching.link_match(db, match, user.id)
    except matching.LinkError as exc:
        raise HTTPException(status.HTTP_409_CONFLICT, str(exc))
    log_action(db, "master.link", company_id=user.company_id, actor_user_id=user.id,
               target_type="master_match", target_id=str(match.id))
    db.commit()
    page = db.get(Page, match.page_id)
    publish_event(user.company_id, "matches.changed", {"action": "linked"})
    publish_event(user.company_id, "corrections.changed", {"action": "master_link"})
    if page:
        publish_event(
            user.company_id, f"document.{page.document_id}.status", {"status": "updated"}
        )
    return {"status": "linked"}


@match_router.post("/matches/{match_id}/dismiss")
def dismiss(
    match_id: uuid.UUID,
    user: User = Depends(require_company_member),
    db: Session = Depends(get_db),
):
    match = _get_match(db, user, match_id)
    matching.dismiss_match(db, match)
    publish_event(user.company_id, "matches.changed", {"action": "dismissed"})
    return {"status": "dismissed"}


@match_router.post("/documents/{document_id}/rematch", status_code=status.HTTP_202_ACCEPTED)
def rematch_document(
    document_id: uuid.UUID,
    user: User = Depends(require_company_member),
    db: Session = Depends(get_db),
):
    document = db.get(Document, document_id)
    if document is None or document.company_id != user.company_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Document not found")
    from app.tasks.ocr_tasks import match_masters

    match_masters.apply_async(args=[str(document.id)])
    return {"status": "queued"}
